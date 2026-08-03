import tomllib
from io import BytesIO
from pathlib import Path
import sys
import types

import pytest

from app.services.rag_document_ingestion_service import RagDocumentParseError, chunk_rag_document


def test_document_parser_stack_is_a_default_backend_dependency() -> None:
    pyproject_path = Path(__file__).resolve().parents[1] / "pyproject.toml"
    pyproject = tomllib.loads(pyproject_path.read_text(encoding="utf-8"))

    dependencies = pyproject["project"]["dependencies"]
    optional_dependencies = pyproject["project"].get("optional-dependencies", {})

    assert "unstructured>=0.18.0,<1.0.0" in dependencies
    assert "python-docx>=1.2.0,<2.0.0" in dependencies
    assert "python-pptx>=1.0.0,<2.0.0" in dependencies
    assert "pdfminer-six>=20251230,<20270000" in dependencies
    assert "pypdf>=6.14.2,<7.0.0" in dependencies
    assert not any(dependency.startswith("unstructured[all-docs]") for dependency in dependencies)
    assert "documents" not in optional_dependencies


def test_backend_dependencies_keep_audited_security_floors() -> None:
    pyproject_path = Path(__file__).resolve().parents[1] / "pyproject.toml"
    pyproject = tomllib.loads(pyproject_path.read_text(encoding="utf-8"))

    dependencies = pyproject["project"]["dependencies"]

    assert "pydantic-settings>=2.14.2,<3.0.0" in dependencies
    assert "starlette>=1.3.1,<2.0.0" in dependencies
    assert "pillow>=12.3.0,<13.0.0" in dependencies
    assert "aiohttp>=3.14.1,<4.0.0" in dependencies
    assert "cryptography>=48.0.1,<49.0.0" in dependencies
    assert "urllib3>=2.7.0,<3.0.0" in dependencies


def test_markdown_document_chunking_preserves_sections_and_overlap() -> None:
    content = """
# 急性腹痛教学资料

急性腹痛训练应先建立病史框架，再决定体格检查与辅助检查。

## 病史采集

需要追问起病时间、疼痛部位、疼痛性质、疼痛程度、疼痛演变和伴随症状。
疼痛迁移、恶心呕吐、发热和既往腹部手术史会影响下一步推理。

## 查体策略

腹部查体应包含视诊、听诊、触诊和腹膜刺激征，避免只做单个压痛点。
如果病史仍不完整，应先补齐关键病史，再解释为什么进入查体。

## 检查选择

血常规、炎症指标、尿常规和腹部超声用于支持或修正诊断假设。
检查结果只能作为反馈和教学知识，不进入标准评分裁判。
""".strip()

    chunks = chunk_rag_document(
        file_name="appendicitis_teaching.md",
        content_bytes=content.encode("utf-8"),
        document_id="kbdoc:appendicitis_001:test",
        max_chars=90,
        overlap_chars=18,
    )

    assert len(chunks) >= 4
    assert [chunk.chunk_index for chunk in chunks] == list(range(len(chunks)))
    assert {chunk.section_title for chunk in chunks} >= {"急性腹痛教学资料", "病史采集", "查体策略", "检查选择"}
    assert all(chunk.source_location.startswith("appendicitis_teaching.md") for chunk in chunks)
    assert all(chunk.document_id == "kbdoc:appendicitis_001:test" for chunk in chunks)
    assert all(chunk.text.strip() for chunk in chunks)
    assert chunks[1].text[:18] in chunks[0].text


def test_unstructured_fallback_chunking_uses_title_strategy_and_adds_chunk_metadata(monkeypatch) -> None:
    captured_partition: dict[str, str] = {}
    captured_chunk_kwargs: dict[str, object] = {}

    class FakeMetadata:
        def __init__(self, page_number: int | None = None) -> None:
            self.page_number = page_number

    class FakeElement:
        def __init__(self, text: str, *, category: str, page_number: int | None) -> None:
            self.text = text
            self.category = category
            self.metadata = FakeMetadata(page_number)

        def __str__(self) -> str:
            return self.text

    class FakeChunk:
        def __init__(self, text: str, orig_elements: list[FakeElement]) -> None:
            self.text = text
            self.category = "CompositeElement"
            self.metadata = types.SimpleNamespace(
                orig_elements=orig_elements,
                page_number=orig_elements[0].metadata.page_number if orig_elements else None,
            )

        def __str__(self) -> str:
            return self.text

    title = FakeElement("Clinical Features", category="Title", page_number=2)
    body = FakeElement("Ask about pain migration and associated fever before deciding next steps.", category="NarrativeText", page_number=2)
    references = FakeElement("References", category="Title", page_number=12)
    reference_body = FakeElement("Smith J. Example article. 2024.", category="NarrativeText", page_number=12)

    def fake_partition(*, filename: str, **kwargs):
        captured_partition["filename"] = filename
        captured_partition["strategy"] = kwargs.get("strategy", "")
        return [title, body, references, reference_body]

    def fake_chunk_by_title(elements, **kwargs):
        captured_chunk_kwargs.update(kwargs)
        assert list(elements) == [title, body, references, reference_body]
        return [
            FakeChunk("Clinical Features\n\nAsk about pain migration and associated fever before deciding next steps.", [title, body]),
            FakeChunk("References\n\nSmith J. Example article. 2024.", [references, reference_body]),
        ]

    partition_module = types.ModuleType("unstructured.partition.auto")
    partition_module.partition = fake_partition
    chunking_module = types.ModuleType("unstructured.chunking.title")
    chunking_module.chunk_by_title = fake_chunk_by_title
    monkeypatch.setitem(sys.modules, "unstructured", types.ModuleType("unstructured"))
    monkeypatch.setitem(sys.modules, "unstructured.partition", types.ModuleType("unstructured.partition"))
    monkeypatch.setitem(sys.modules, "unstructured.partition.auto", partition_module)
    monkeypatch.setitem(sys.modules, "unstructured.chunking", types.ModuleType("unstructured.chunking"))
    monkeypatch.setitem(sys.modules, "unstructured.chunking.title", chunking_module)

    chunks = chunk_rag_document(
        file_name="appendicitis_review.odt",
        content_bytes=b"pretend-open-document-content",
        document_id="kbdoc:appendicitis_001:review",
        max_chars=900,
        overlap_chars=120,
    )

    assert captured_partition["filename"].endswith(".odt")
    assert captured_partition["strategy"] == ""
    assert captured_chunk_kwargs["max_characters"] == 900
    assert captured_chunk_kwargs["overlap"] == 120
    assert captured_chunk_kwargs["include_orig_elements"] is True
    assert captured_chunk_kwargs["multipage_sections"] is True

    assert len(chunks) == 2
    assert chunks[0].chunking_strategy == "unstructured_by_title"
    assert chunks[0].chunk_categories == ["Title", "NarrativeText"]
    assert chunks[0].section_title == "Clinical Features"
    assert chunks[0].page_number == 2
    assert chunks[0].char_count == len(chunks[0].text)
    assert chunks[0].source_location == "appendicitis_review.odt · 第 2 页 · Clinical Features · 片段 1"
    assert chunks[1].risk_flags == ["references_section"]
    assert "low_value_section" in chunks[1].quality_warnings


def test_pdf_document_uses_local_pypdf_text_extraction(monkeypatch) -> None:
    import pypdf

    class FakePage:
        def __init__(self, text: str) -> None:
            self._text = text

        def extract_text(self) -> str:
            return self._text

    class FakeReader:
        pages = [
            FakePage("腹痛问诊\n\n应追问起病部位、疼痛迁移和伴随症状。"),
            FakePage("查体策略\n\n腹部查体应覆盖视诊、听诊、触诊和腹膜刺激征。"),
        ]

    monkeypatch.setattr(pypdf, "PdfReader", lambda stream: FakeReader())

    chunks = chunk_rag_document(
        file_name="acute_abdomen.pdf",
        content_bytes=b"%PDF-local-parser-test",
        document_id="kbdoc:appendicitis_001:pdf",
    )

    assert [chunk.page_number for chunk in chunks] == [1, 2]
    assert [chunk.section_title for chunk in chunks] == ["腹痛问诊", "查体策略"]
    assert all(chunk.chunking_strategy == "local_pdf_page_window" for chunk in chunks)


def test_scanned_pdf_without_selectable_text_requests_ocr(monkeypatch) -> None:
    import pypdf

    class FakePage:
        def extract_text(self) -> str:
            return ""

    monkeypatch.setattr(pypdf, "PdfReader", lambda stream: types.SimpleNamespace(pages=[FakePage()]))

    with pytest.raises(RagDocumentParseError, match="require OCR"):
        chunk_rag_document(
            file_name="scanned.pdf",
            content_bytes=b"%PDF-scanned-test",
            document_id="kbdoc:global:scanned",
        )


def test_docx_document_uses_local_section_and_table_extraction() -> None:
    from docx import Document

    document = Document()
    document.add_heading("腹痛问诊", level=1)
    document.add_paragraph("先追问起病、部位、性质、程度、迁移和伴随症状。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "阶段"
    table.cell(0, 1).text = "训练目标"
    stream = BytesIO()
    document.save(stream)

    chunks = chunk_rag_document(
        file_name="acute_abdomen.docx",
        content_bytes=stream.getvalue(),
        document_id="kbdoc:appendicitis_001:docx",
    )

    assert len(chunks) == 1
    assert chunks[0].section_title == "腹痛问诊"
    assert chunks[0].chunking_strategy == "local_docx_section_window"
    assert "Table" in chunks[0].chunk_categories
    assert "阶段 | 训练目标" in chunks[0].text


def test_pptx_document_uses_local_slide_extraction() -> None:
    from pptx import Presentation

    presentation = Presentation()
    first_slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    first_slide.shapes.title.text = "病史采集"
    first_slide.placeholders[1].text = "追问疼痛迁移、恶心呕吐和发热。"
    second_slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    second_slide.shapes.title.text = "辅助检查"
    second_slide.placeholders[1].text = "根据当前诊断假设选择检查并说明目的。"
    stream = BytesIO()
    presentation.save(stream)

    chunks = chunk_rag_document(
        file_name="osce_teaching.pptx",
        content_bytes=stream.getvalue(),
        document_id="kbdoc:appendicitis_001:pptx",
    )

    assert [chunk.page_number for chunk in chunks] == [1, 2]
    assert [chunk.section_title for chunk in chunks] == ["病史采集", "辅助检查"]
    assert all(chunk.chunking_strategy == "local_pptx_slide_window" for chunk in chunks)
    assert "追问疼痛迁移" in chunks[0].text


def test_html_document_uses_local_readable_content_extraction() -> None:
    chunks = chunk_rag_document(
        file_name="teacher_note.html",
        content_bytes=(
            "<html><head><style>.hidden{display:none}</style><script>secret()</script></head>"
            "<body><h1>沟通训练</h1><p>先开放式询问，再总结确认患者的担忧。</p></body></html>"
        ).encode("utf-8"),
        document_id="kbdoc:global:html",
    )

    assert len(chunks) == 1
    assert chunks[0].section_title == "沟通训练"
    assert chunks[0].chunking_strategy == "local_html_section_window"
    assert "开放式询问" in chunks[0].text
    assert "secret" not in chunks[0].text
